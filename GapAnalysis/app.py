import io
import datetime
import pandas as pd
import numpy as np
from flask import Flask, render_template, request, jsonify, send_file, send_from_directory
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import (
    PRIMARY_TICKERS,
    POPULAR_TICKERS,
    DEFAULT_TARGET_MINUTES,
    DEFAULT_GAP_THRESHOLD_PCT
)
from db_loader import (
    get_available_tickers,
    get_data_source_status,
    prefetch_primary_tickers
)
from gap_engine import (
    analyze_ticker_gaps,
    compare_multiple_tickers,
    analyze_premarket_momentum,
    compare_premarket_tickers
)

app = Flask(
    __name__,
    static_folder="static",
    template_folder="templates"
)


# ============================================================
# UTILITIES
# ============================================================

def clean_json_data(obj):
    """Recursively clean NaNs, Infinities and NumPy datatypes for JSON serialization."""
    if obj is None:
        return None
    if isinstance(obj, (np.integer, int)):
        return int(obj)
    if isinstance(obj, (np.floating, float)):
        if np.isnan(obj) or not np.isfinite(obj):
            return None
        return float(obj)
    if isinstance(obj, (np.ndarray, list)):
        return [clean_json_data(x) for x in obj]
    if isinstance(obj, dict):
        return {str(k): clean_json_data(v) for k, v in obj.items()}
    return obj


# ============================================================
# ROUTES
# ============================================================

@app.route("/")
def index():
    return send_from_directory(app.static_folder, "index.html")


@app.route("/api/status", methods=["GET"])
def api_status():
    status = get_data_source_status()
    return jsonify(clean_json_data(status))


@app.route("/api/tickers", methods=["GET"])
def api_tickers():
    tickers = get_available_tickers()
    return jsonify({"tickers": tickers, "primary": PRIMARY_TICKERS})


@app.route("/api/premarket", methods=["POST"])
def api_premarket():
    try:
        data = request.get_json() or {}
        ticker = str(data.get("ticker", "QQQ")).strip().upper()
        days_back = int(data.get("days_back", 365))
        threshold_pct = float(data.get("threshold_pct", 0.1))

        res = analyze_premarket_momentum(
            ticker=ticker,
            days_back=days_back,
            threshold_pct=threshold_pct
        )

        if "error" in res:
            return jsonify({"error": res["error"]}), 400

        return jsonify(clean_json_data(res))
    except Exception as e:
        return jsonify({"error": f"Ошибка премаркет анализа: {str(e)}"}), 500


@app.route("/api/premarket-compare", methods=["POST"])
def api_premarket_compare():
    try:
        data = request.get_json() or {}
        tickers = data.get("tickers", PRIMARY_TICKERS)
        if not isinstance(tickers, list) or not tickers:
            tickers = PRIMARY_TICKERS

        res = compare_premarket_tickers(tickers=tickers)
        return jsonify(clean_json_data({"comparison": res}))
    except Exception as e:
        return jsonify({"error": f"Ошибка сравнения премаркета: {str(e)}"}), 500


@app.route("/api/analyze", methods=["POST"])
def api_analyze():
    try:
        data = request.get_json() or {}
        ticker = str(data.get("ticker", "QQQ")).strip().upper()
        target_minutes = int(data.get("target_minutes", DEFAULT_TARGET_MINUTES))
        gap_threshold_pct = float(data.get("gap_threshold_pct", DEFAULT_GAP_THRESHOLD_PCT))
        days_back = int(data.get("days_back", 365))
        start_date = data.get("start_date")
        end_date = data.get("end_date")

        res = analyze_ticker_gaps(
            ticker=ticker,
            target_minutes=target_minutes,
            gap_threshold_pct=gap_threshold_pct,
            days_back=days_back,
            start_date=start_date,
            end_date=end_date
        )

        if "error" in res:
            return jsonify({"error": res["error"]}), 400

        return jsonify(clean_json_data(res))
    except Exception as e:
        return jsonify({"error": f"Ошибка анализа: {str(e)}"}), 500


@app.route("/api/compare", methods=["POST"])
def api_compare():
    try:
        data = request.get_json() or {}
        tickers = data.get("tickers", PRIMARY_TICKERS)
        if not isinstance(tickers, list) or not tickers:
            tickers = PRIMARY_TICKERS

        target_minutes = int(data.get("target_minutes", DEFAULT_TARGET_MINUTES))
        gap_threshold_pct = float(data.get("gap_threshold_pct", DEFAULT_GAP_THRESHOLD_PCT))
        days_back = int(data.get("days_back", 365))

        res = compare_multiple_tickers(
            tickers=tickers,
            target_minutes=target_minutes,
            gap_threshold_pct=gap_threshold_pct,
            days_back=days_back
        )

        return jsonify(clean_json_data({"comparison": res}))
    except Exception as e:
        return jsonify({"error": f"Ошибка сравнения: {str(e)}"}), 500


@app.route("/api/export", methods=["POST"])
def api_export():
    """Generates a professional multi-tab Excel report."""
    try:
        data = request.get_json() or {}
        ticker = str(data.get("ticker", "QQQ")).strip().upper()
        target_minutes = int(data.get("target_minutes", DEFAULT_TARGET_MINUTES))
        gap_threshold_pct = float(data.get("gap_threshold_pct", DEFAULT_GAP_THRESHOLD_PCT))
        days_back = int(data.get("days_back", 365))

        res = analyze_ticker_gaps(
            ticker=ticker,
            target_minutes=target_minutes,
            gap_threshold_pct=gap_threshold_pct,
            days_back=days_back
        )

        if "error" in res:
            return jsonify({"error": res["error"]}), 400

        records = res.get("daily_records", [])
        summary = res.get("summary", {})
        up_stats = summary.get("gap_up_stats", {})
        down_stats = summary.get("gap_down_stats", {})

        wb = openpyxl.Workbook()
        ws_summary = wb.active
        ws_summary.title = "Сводка (Summary)"

        # Style definitions
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        title_font = Font(name="Arial", size=14, bold=True, color="0F172A")
        bold_font = Font(name="Arial", size=10, bold=True)
        regular_font = Font(name="Arial", size=10)
        up_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
        down_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E1"),
            right=Side(style="thin", color="CBD5E1"),
            top=Side(style="thin", color="CBD5E1"),
            bottom=Side(style="thin", color="CBD5E1")
        )

        # Tab 1: Summary Info
        ws_summary.append([f"Анализ Гэпов и Динамики Цены: {ticker}"])
        ws_summary.cell(row=1, column=1).font = title_font
        ws_summary.append([f"Срез времени: {target_minutes} мин после открытия ({res['parameters']['target_time_et']}) | Порог гэпа: >={gap_threshold_pct}% | Дней: {summary.get('total_days')}"])
        ws_summary.append([])

        summary_rows = [
            ["Метрика", "Gap Up (Гэп Вверх)", "Gap Down (Гэп Вниз)"],
            ["Количество торговых дней", up_stats.get("count"), down_stats.get("count")],
            ["Доля от всех дней (%)", f"{up_stats.get('pct_of_total')}%", f"{down_stats.get('pct_of_total')}%"],
            ["Средний гэп при открытии (%)", f"{up_stats.get('avg_gap_pct')}%", f"{down_stats.get('avg_gap_pct')}%"],
            ["Вероятность продолжения (Win Rate %)", f"{up_stats.get('continuation_rate_pct')}%", f"{down_stats.get('continuation_rate_pct')}%"],
            ["Вероятность разворота (%)", f"{up_stats.get('reversal_rate_pct')}%", f"{down_stats.get('reversal_rate_pct')}%"],
            ["Средний % от Prev. Close к моменту T", f"{up_stats.get('avg_move_from_prev_close_pct')}%", f"{down_stats.get('avg_move_from_prev_close_pct')}%"],
            ["Медианный % от Prev. Close к моменту T", f"{up_stats.get('median_move_from_prev_close_pct')}%", f"{down_stats.get('median_move_from_prev_close_pct')}%"],
            ["Средний % от Open к моменту T", f"{up_stats.get('avg_move_from_open_pct')}%", f"{down_stats.get('avg_move_from_open_pct')}%"],
            ["Средний макс. импульс в сторону гэпа (%)", f"{up_stats.get('avg_favorable_move_pct')}%", f"{down_stats.get('avg_favorable_move_pct')}%"],
            ["Средняя просадка против гэпа (%)", f"{up_stats.get('avg_adverse_move_pct')}%", f"{down_stats.get('avg_adverse_move_pct')}%"],
            ["Частота закрытия гэпа в интервале (%)", f"{up_stats.get('gap_fill_rate_pct')}%", f"{down_stats.get('gap_fill_rate_pct')}%"],
            ["Средний результат полного дня (%)", f"{up_stats.get('avg_full_day_pct')}%", f"{down_stats.get('avg_full_day_pct')}%"]
        ]

        for r_idx, row_vals in enumerate(summary_rows, start=4):
            for c_idx, val in enumerate(row_vals, start=1):
                cell = ws_summary.cell(row=r_idx, column=c_idx, value=val)
                cell.border = thin_border
                if r_idx == 4:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.font = bold_font if c_idx == 1 else regular_font
                    cell.alignment = Alignment(horizontal="left" if c_idx == 1 else "center", vertical="center")

        # Tab 2: Detailed Daily Data
        ws_data = wb.create_sheet(title="История по дням (Daily)")
        data_headers = [
            "Дата", "Тикер", "Prev Close ($)", "Open ($)", "Цена в T ($)",
            "High дня ($)", "Low дня ($)", "Close дня ($)", "Гэп ($)", "Гэп (%)",
            "Тип гэпа", "Изм. от Prev Close (%)", "Изм. от Open (%)",
            "Продолжение?", "Гэп закрыт?", "Макс. импульс (%)", "Макс. просадка (%)"
        ]

        ws_data.append(data_headers)
        for col_idx in range(1, len(data_headers) + 1):
            cell = ws_data.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

        for r_idx, r in enumerate(records, start=2):
            row_data = [
                r["date"], r["ticker"], r["prev_close"], r["open"], r["target_price"],
                r["day_high"], r["day_low"], r["day_close"], r["gap_dollar"], r["gap_pct"],
                r["gap_type"], r["move_from_prev_close_pct"], r["move_from_open_pct"],
                "ДА" if r["continuation"] else "НЕТ",
                "ДА" if r["gap_filled"] else "НЕТ",
                r["favorable_move_pct"], r["adverse_move_pct"]
            ]
            for c_idx, val in enumerate(row_data, start=1):
                cell = ws_data.cell(row=r_idx, column=c_idx, value=val)
                cell.font = regular_font
                cell.border = thin_border
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if r["gap_type"] == "GAP_UP":
                    cell.fill = up_fill
                elif r["gap_type"] == "GAP_DOWN":
                    cell.fill = down_fill

        # Auto-adjust column widths for all sheets
        for sheet in wb.worksheets:
            for col in sheet.columns:
                max_len = 0
                col_letter = get_column_letter(col[0].column)
                for cell in col:
                    if cell.value:
                        max_len = max(max_len, len(str(cell.value)))
                sheet.column_dimensions[col_letter].width = max(max_len + 4, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"gap_analysis_{ticker}_{datetime.date.today()}.xlsx"
        return send_file(
            output,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            as_attachment=True,
            download_name=filename
        )
    except Exception as e:
        return jsonify({"error": f"Ошибка формирования Excel: {str(e)}"}), 500


# Warmup cache on startup
try:
    prefetch_primary_tickers()
except Exception as e:
    print("Warmup error:", e)


if __name__ == "__main__":
    print("Starting Gap Analysis Server on http://localhost:5000")
    app.run(host="0.0.0.0", port=5000, debug=True)
