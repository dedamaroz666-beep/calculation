from flask import (
    Flask,
    render_template,
    request,
    jsonify,
    send_file
)

import numpy as np
import io
import datetime
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import BASE_DIR, BENCHMARKS
from database import (
    get_stock_data,
    get_available_tickers,
    get_database_status,
    get_intraday_chart_data,
    get_ticker_reports,
    get_ticker_dividends_db,
    save_ticker_dividends_db,
    get_company_profile_db,
    save_company_profile_db
)

from calculations import (
    calculate_all_periods,
    calculate_groups_analysis
)
from services.polygon_service import polygon_client
from services.fmp_service import fmp_client
from services.earnings_reaction_service import earnings_reaction_client
from services.quant_service import quant_client
from services.earnings_extra_service import get_ticker_earnings_extra_map, find_matching_earnings_extra


app = Flask(__name__)


# ============================================================
# HELPERS
# ============================================================

def clean_value(value):
    if value is None:
        return None

    if isinstance(value, (np.integer, int)):
        return int(value)

    if isinstance(value, (np.floating, float)):
        if np.isnan(value) or not np.isfinite(value):
            return None
        return float(value)

    return value


def clean_dict(data):
    if isinstance(data, dict):
        return {
            str(key): clean_dict(value)
            for key, value in data.items()
        }

    if isinstance(data, list):
        return [
            clean_dict(value)
            for value in data
        ]

    return clean_value(data)


# ============================================================
# MAIN PAGE
# ============================================================

@app.route("/")
def index():
    try:
        tickers = get_available_tickers()
    except Exception as e:
        print("Database error:", e)
        tickers = []

    return render_template(
        "index.html",
        tickers=tickers
    )


# ============================================================
# SINGLE ASSET CALCULATION
# ============================================================

@app.route("/api/calculate", methods=["POST"])
def calculate():
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "Пустой запрос."}), 400

        stock = str(data.get("stock", "")).strip().upper()
        benchmark = str(data.get("benchmark", "")).strip().upper()
        benchmark_move = data.get("benchmark_move", 0.01)
        asset_move = data.get("asset_move")

        try:
            benchmark_move = float(benchmark_move)
        except (TypeError, ValueError):
            return jsonify({"error": "Некорректное движение benchmark."}), 400

        if asset_move is not None:
            try:
                asset_move = float(asset_move)
            except (TypeError, ValueError):
                return jsonify({"error": "Некорректное движение акции."}), 400

        if not stock:
            return jsonify({"error": "Не указана акция."}), 400

        if not benchmark:
            return jsonify({"error": "Не указан benchmark."}), 400

        if stock == benchmark:
            return jsonify({"error": "Акции должны быть разными."}), 400

        df = get_stock_data(stock, benchmark)
        if df is None or df.empty:
            return jsonify({"error": "Данные не найдены."}), 404

        results, _ = calculate_all_periods(
            df,
            stock,
            benchmark,
            benchmark_move=benchmark_move,
            asset_move=asset_move
        )

        return jsonify({
            "stock": stock,
            "benchmark": benchmark,
            "benchmark_move": benchmark_move,
            "asset_move": asset_move,
            "results": clean_dict(results),
            "chart": []
        })

    except Exception as e:
        print("Calculation error:", e)
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# ============================================================
# GROUP / BATCH ANALYSIS (UP TO 300+ TICKERS)
# ============================================================

# ============================================================
# GROUP / BATCH ANALYSIS (UP TO 300+ TICKERS & MULTI-BENCHMARK)
# ============================================================

@app.route("/api/calculate-groups", methods=["POST"])
@app.route("/api/calculate_groups", methods=["POST"])
def calculate_groups():
    import re
    try:
        data = request.get_json()
        if not data:
            return jsonify({"error": "Пустой запрос."}), 400

        groups = data.get("groups")
        benchmark_raw = data.get("benchmarks") or data.get("benchmark") or data.get("reference")
        period = data.get("period") or "1y"
        benchmark_move = data.get("benchmark_move", 0.01)
        benchmark_moves = data.get("benchmark_moves", {})
        asset_moves = data.get("asset_moves", {})

        try:
            benchmark_move = float(benchmark_move)
        except (TypeError, ValueError):
            return jsonify({"error": "Некорректное движение benchmark."}), 400

        if not groups or not isinstance(groups, dict):
            return jsonify({"error": "Добавьте хотя бы одну группу с тикерами."}), 400

        if not benchmark_raw:
            return jsonify({"error": "Не указан Market Asset (Benchmark)."}), 400

        # Parse multiple benchmarks
        benchmarks_list = []
        if isinstance(benchmark_raw, (list, tuple, set)):
            benchmarks_list = [str(b).strip().upper() for b in benchmark_raw if str(b).strip()]
        elif isinstance(benchmark_raw, str):
            benchmarks_list = [b.strip().upper() for b in re.split(r"[\s,;]+", benchmark_raw) if b.strip()]

        clean_benchmarks = []
        for bm in benchmarks_list:
            if bm and bm not in clean_benchmarks:
                clean_benchmarks.append(bm)

        if not clean_benchmarks:
            return jsonify({"error": "Не указан валидный Market Asset (Benchmark)."}), 400

        # Clean Groups & Tickers
        cleaned_groups = {}
        all_requested_tickers = []

        for group_name, tickers in groups.items():
            if not isinstance(tickers, list):
                continue

            clean_tickers = []
            for ticker in tickers:
                if not isinstance(ticker, str):
                    continue
                ticker = ticker.strip().upper()
                if ticker and ticker not in clean_benchmarks and ticker not in clean_tickers:
                    clean_tickers.append(ticker)
                    if ticker not in all_requested_tickers:
                        all_requested_tickers.append(ticker)

            if clean_tickers:
                name = str(group_name).strip() or "Group"
                cleaned_groups[name] = clean_tickers

        if not cleaned_groups or not all_requested_tickers:
            return jsonify({"error": "Нет валидных тикеров для анализа."}), 400

        all_symbols = list(dict.fromkeys(all_requested_tickers + clean_benchmarks))

        # Fetch Data from Database
        df = get_stock_data(*all_symbols)
        if df is None or df.empty:
            return jsonify({"error": "Данные для указанных тикеров не найдены в базе данных."}), 404

        available_tickers = set(
            df["ticker"].astype(str).str.strip().str.upper().unique()
        )

        missing_benchmarks = [bm for bm in clean_benchmarks if bm not in available_tickers]
        if missing_benchmarks:
            return jsonify({
                "error": f"Market Asset '{', '.join(missing_benchmarks)}' не найден в базе данных. Проверьте правильность тикеров."
            }), 404

        # Check missing tickers (graceful handling for 300+ batch)
        missing_tickers = [
            ticker for ticker in all_requested_tickers
            if ticker not in available_tickers
        ]

        # Filter groups to only available tickers
        active_groups = {}
        for g_name, t_list in cleaned_groups.items():
            valid_t = [t for t in t_list if t in available_tickers]
            if valid_t:
                active_groups[g_name] = valid_t

        if not active_groups:
            return jsonify({
                "error": "Ни один из указанных тикеров не найден в базе данных.",
                "missing": missing_tickers
            }), 404

        # Clean Asset Moves
        if not isinstance(asset_moves, dict):
            asset_moves = {}

        cleaned_asset_moves = {}
        for ticker, move in asset_moves.items():
            ticker = str(ticker).strip().upper()
            try:
                move = float(move)
            except (TypeError, ValueError):
                continue
            if np.isfinite(move):
                cleaned_asset_moves[ticker] = move

        # Run Vectorized Multi-Benchmark Group Analysis
        analysis = calculate_groups_analysis(
            df=df,
            groups=active_groups,
            benchmarks=clean_benchmarks,
            period=period,
            benchmark_move=benchmark_move,
            benchmark_moves=benchmark_moves if isinstance(benchmark_moves, dict) else None,
            asset_moves=cleaned_asset_moves
        )

        ignored_tickers = analysis.get("ignored_tickers", {})
        partial_tickers = analysis.get("partial_tickers", {})
        warning_parts = []
        if missing_tickers:
            warning_parts.append(
                f"Пропущено {len(missing_tickers)} тикеров (нет в БД): {', '.join(missing_tickers[:8])}"
                + (f" и ещё {len(missing_tickers) - 8}..." if len(missing_tickers) > 8 else "")
            )
        if ignored_tickers:
            ignored_list = list(ignored_tickers.keys()) if isinstance(ignored_tickers, dict) else list(ignored_tickers)
            warning_parts.append(
                f"Не удалось рассчитать {len(ignored_list)} тикеров (< 2 наблюдений): {', '.join(ignored_list[:8])}"
                + (f" и ещё {len(ignored_list) - 8}..." if len(ignored_list) > 8 else "")
            )
        if partial_tickers:
            partial_list = list(partial_tickers.keys())
            warning_parts.append(
                f"⚠️ {len(partial_list)} тикеров имеют историю меньше периода {period} и подсвечены в таблице: {', '.join(partial_list[:6])}"
                + (f" и ещё {len(partial_list) - 6}..." if len(partial_list) > 6 else "")
            )
        warning_message = " | ".join(warning_parts) if warning_parts else None

        response = {
            "groups": active_groups,
            "benchmarks": clean_benchmarks,
            "primary_benchmark": clean_benchmarks[0],
            "benchmark": clean_benchmarks[0],
            "reference": clean_benchmarks[0],
            "benchmark_moves": clean_dict(analysis.get("benchmark_moves", {})),
            "benchmark_move": benchmark_move,
            "asset_moves": cleaned_asset_moves,
            "tickers": analysis.get("tickers", []),
            "all_symbols": analysis.get("all_symbols", []),
            "period": analysis.get("period", period),
            "period_start": analysis.get("period_start"),
            "period_end": analysis.get("period_end"),
            "correlation": clean_dict(analysis.get("correlation", {})),
            "beta": clean_dict(analysis.get("beta", {})),
            "benchmark_beta": clean_dict(analysis.get("benchmark_beta", {})),
            "expected_range": clean_dict(analysis.get("expected_range", {})),
            "benchmarks_data": clean_dict(analysis.get("benchmarks_data", {})),
            "group_summary": clean_dict(analysis.get("group_summary", {})),
            "missing_tickers": missing_tickers,
            "ignored_tickers": clean_dict(ignored_tickers),
            "partial_tickers": clean_dict(partial_tickers),
            "warning": warning_message,
            "total_requested": len(all_requested_tickers),
            "total_analyzed": len(analysis.get("tickers", []))
        }

        return jsonify(response)

    except Exception as e:
        print("Group calculation error:", e)
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# ============================================================
# EXPORT FILTERED CORRELATION RESULTS TO EXCEL (.XLSX)
# ============================================================

@app.route("/api/export-excel", methods=["POST"])
def export_excel():
    try:
        data = request.get_json() or {}

        rows = data.get("rows", [])
        benchmarks = data.get("benchmarks", [])
        benchmark = str(data.get("benchmark", "")).strip().upper() or "BENCHMARK"
        benchmarks_data = data.get("benchmarks_data", {})
        period = str(data.get("period", "1y"))
        filter_label = str(data.get("filter_label", "All Assets")).strip()
        is_multi_benchmark = bool(data.get("is_multi_benchmark", False) or (isinstance(benchmarks, list) and len(benchmarks) > 1))

        if not isinstance(rows, list) or not rows:
            return jsonify({"error": "Нет результатов для экспорта в Excel."}), 400

        if not benchmarks:
            benchmarks = [benchmark]

        wb = openpyxl.Workbook()

        # Styles definition
        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_fill_accent = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        border_thin = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        # Correlation fills
        fill_high_pos = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid") # soft green
        fill_mod_pos = PatternFill(start_color="F0FDF4", end_color="F0FDF4", fill_type="solid")  # light green
        fill_high_neg = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid") # soft red
        fill_mod_neg = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")  # light red

        # Format helper for correlation cells
        def format_corr_cell(cell, corr):
            if corr is not None:
                cell.number_format = "0.000"
                if corr >= 0.50:
                    cell.fill = fill_high_pos
                    cell.font = Font(name="Calibri", size=10, bold=True, color="166534")
                elif corr >= 0.20:
                    cell.fill = fill_mod_pos
                    cell.font = Font(name="Calibri", size=10, color="15803D")
                elif corr <= -0.50:
                    cell.fill = fill_high_neg
                    cell.font = Font(name="Calibri", size=10, bold=True, color="991B1B")
                elif corr < 0:
                    cell.fill = fill_mod_neg
                    cell.font = Font(name="Calibri", size=10, color="B91C1C")

        # ========================================================
        # SHEET 1: Main View (Multi-Benchmark Comparison OR Single Benchmark)
        # ========================================================
        ws1 = wb.active

        if is_multi_benchmark and len(benchmarks) > 1:
            ws1.title = "Benchmark Comparison"

            headers = ["Ticker"]
            for bm in benchmarks:
                headers.extend([f"Corr ({bm})", f"Beta ({bm})", f"Exp Move ({bm})"])
            headers.extend(["Actual Move", "Daily Volatility", "Observations", "Period"])

            ws1.append(headers)
            ws1.row_dimensions[1].height = 26

            for col_num in range(1, len(headers) + 1):
                cell = ws1.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = cell_border

            # Populate Rows
            for row_idx, r in enumerate(rows, start=2):
                ticker = str(r.get("ticker", "")).strip().upper()
                row_data = [ticker]

                for bm in benchmarks:
                    bm_corr = clean_value(r.get(f"correlation_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("correlation"))
                    bm_beta = clean_value(r.get(f"beta_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("beta"))
                    bm_exp = clean_value(r.get(f"expected_move_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("expected_move"))
                    row_data.extend([bm_corr, bm_beta, bm_exp])

                act_move = clean_value(r.get("actual_move"))
                vol = clean_value(r.get("volatility"))
                obs = clean_value(r.get("observations"))

                row_data.extend([act_move, vol, obs, period])
                ws1.append(row_data)
                ws1.row_dimensions[row_idx].height = 20

                # Format Cells
                col_c = 1
                c_ticker = ws1.cell(row=row_idx, column=col_c)
                c_ticker.alignment = Alignment(horizontal="center", vertical="center")
                c_ticker.font = Font(name="Calibri", size=10, bold=True)
                c_ticker.border = cell_border
                col_c += 1

                for bm in benchmarks:
                    # Corr
                    c_corr = ws1.cell(row=row_idx, column=col_c)
                    c_corr.alignment = Alignment(horizontal="right", vertical="center")
                    c_corr.border = cell_border
                    bm_corr = clean_value(r.get(f"correlation_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("correlation"))
                    format_corr_cell(c_corr, bm_corr)
                    col_c += 1

                    # Beta
                    c_beta = ws1.cell(row=row_idx, column=col_c)
                    c_beta.alignment = Alignment(horizontal="right", vertical="center")
                    c_beta.border = cell_border
                    bm_beta = clean_value(r.get(f"beta_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("beta"))
                    if bm_beta is not None:
                        c_beta.number_format = "0.000"
                    col_c += 1

                    # Exp Move
                    c_exp = ws1.cell(row=row_idx, column=col_c)
                    c_exp.alignment = Alignment(horizontal="right", vertical="center")
                    c_exp.border = cell_border
                    bm_exp = clean_value(r.get(f"expected_move_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("expected_move"))
                    if bm_exp is not None:
                        c_exp.number_format = "0.00%"
                    col_c += 1

                # Actual move
                c_act = ws1.cell(row=row_idx, column=col_c)
                c_act.alignment = Alignment(horizontal="right", vertical="center")
                c_act.border = cell_border
                if act_move is not None:
                    c_act.number_format = "0.00%"
                col_c += 1

                # Volatility
                c_vol = ws1.cell(row=row_idx, column=col_c)
                c_vol.alignment = Alignment(horizontal="right", vertical="center")
                c_vol.border = cell_border
                if vol is not None:
                    c_vol.number_format = "0.00%"
                col_c += 1

                # Obs
                c_obs = ws1.cell(row=row_idx, column=col_c)
                c_obs.alignment = Alignment(horizontal="center", vertical="center")
                c_obs.border = cell_border
                if obs is not None:
                    c_obs.number_format = "#,##0"
                col_c += 1

                # Period
                c_per = ws1.cell(row=row_idx, column=col_c)
                c_per.alignment = Alignment(horizontal="center", vertical="center")
                c_per.border = cell_border

            ws1.freeze_panes = "A2"
            ws1.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"

            # Auto column widths
            for col_idx in range(1, len(headers) + 1):
                col_letter = get_column_letter(col_idx)
                header_text = str(headers[col_idx - 1])
                ws1.column_dimensions[col_letter].width = max(len(header_text) + 4, 13)

            # ========================================================
            # ADDITIONAL SHEETS FOR EACH INDIVIDUAL BENCHMARK
            # ========================================================
            for bm in benchmarks:
                sheet_title = f"{bm} Analysis"[:31]
                ws_bm = wb.create_sheet(title=sheet_title)

                bm_headers = [
                    "Ticker",
                    "Correlation",
                    "Beta",
                    "Expected Move",
                    "Actual Move",
                    "Difference",
                    "Normalized Diff",
                    "Daily Volatility",
                    "Residual Volatility",
                    "Observations",
                    "Benchmark",
                    "Period"
                ]
                ws_bm.append(bm_headers)
                ws_bm.row_dimensions[1].height = 26

                for col_num in range(1, len(bm_headers) + 1):
                    cell = ws_bm.cell(row=1, column=col_num)
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_alignment
                    cell.border = cell_border

                for row_idx, r in enumerate(rows, start=2):
                    ticker = str(r.get("ticker", "")).strip().upper()
                    m = r.get("benchmarks_metrics", {}).get(bm, {}) if isinstance(r.get("benchmarks_metrics"), dict) else {}
                    corr = clean_value(r.get(f"correlation_{bm}", m.get("correlation", r.get("correlation"))))
                    beta_val = clean_value(r.get(f"beta_{bm}", m.get("beta", r.get("beta"))))
                    exp_move = clean_value(r.get(f"expected_move_{bm}", m.get("expected_move", r.get("expected_move"))))
                    act_move = clean_value(r.get("actual_move", m.get("actual_move")))
                    diff = clean_value(m.get("difference", r.get(f"difference_{bm}", r.get("difference"))))
                    norm = clean_value(m.get("normalized", r.get(f"normalized_{bm}", r.get("normalized"))))
                    vol = clean_value(r.get("volatility", m.get("volatility")))
                    res_vol = clean_value(m.get("residual_volatility", r.get(f"residual_volatility_{bm}", r.get("residual_volatility"))))
                    obs = clean_value(r.get("observations", m.get("observations")))

                    row_data = [
                        ticker,
                        corr,
                        beta_val,
                        exp_move,
                        act_move,
                        diff,
                        norm,
                        vol,
                        res_vol,
                        obs,
                        bm,
                        period
                    ]
                    ws_bm.append(row_data)
                    ws_bm.row_dimensions[row_idx].height = 20

                    # Styling
                    c_tk = ws_bm.cell(row=row_idx, column=1)
                    c_tk.alignment = Alignment(horizontal="center", vertical="center")
                    c_tk.font = Font(name="Calibri", size=10, bold=True)
                    c_tk.border = cell_border

                    c_cr = ws_bm.cell(row=row_idx, column=2)
                    c_cr.alignment = Alignment(horizontal="right", vertical="center")
                    c_cr.border = cell_border
                    format_corr_cell(c_cr, corr)

                    c_bt = ws_bm.cell(row=row_idx, column=3)
                    c_bt.alignment = Alignment(horizontal="right", vertical="center")
                    c_bt.border = cell_border
                    if beta_val is not None:
                        c_bt.number_format = "0.000"

                    c_ex = ws_bm.cell(row=row_idx, column=4)
                    c_ex.alignment = Alignment(horizontal="right", vertical="center")
                    c_ex.border = cell_border
                    if exp_move is not None:
                        c_ex.number_format = "0.00%"

                    c_ac = ws_bm.cell(row=row_idx, column=5)
                    c_ac.alignment = Alignment(horizontal="right", vertical="center")
                    c_ac.border = cell_border
                    if act_move is not None:
                        c_ac.number_format = "0.00%"

                    c_df = ws_bm.cell(row=row_idx, column=6)
                    c_df.alignment = Alignment(horizontal="right", vertical="center")
                    c_df.border = cell_border
                    if diff is not None:
                        c_df.number_format = "0.00%"
                        if diff > 0:
                            c_df.font = Font(name="Calibri", size=10, color="16A34A")
                        elif diff < 0:
                            c_df.font = Font(name="Calibri", size=10, color="DC2626")

                    c_nm = ws_bm.cell(row=row_idx, column=7)
                    c_nm.alignment = Alignment(horizontal="right", vertical="center")
                    c_nm.border = cell_border
                    if norm is not None:
                        c_nm.number_format = "0.00"

                    c_vl = ws_bm.cell(row=row_idx, column=8)
                    c_vl.alignment = Alignment(horizontal="right", vertical="center")
                    c_vl.border = cell_border
                    if vol is not None:
                        c_vl.number_format = "0.00%"

                    c_rv = ws_bm.cell(row=row_idx, column=9)
                    c_rv.alignment = Alignment(horizontal="right", vertical="center")
                    c_rv.border = cell_border
                    if res_vol is not None:
                        c_rv.number_format = "0.00%"

                    c_ob = ws_bm.cell(row=row_idx, column=10)
                    c_ob.alignment = Alignment(horizontal="center", vertical="center")
                    c_ob.border = cell_border
                    if obs is not None:
                        c_ob.number_format = "#,##0"

                    c_bm_c = ws_bm.cell(row=row_idx, column=11)
                    c_bm_c.alignment = Alignment(horizontal="center", vertical="center")
                    c_bm_c.border = cell_border

                    c_pr = ws_bm.cell(row=row_idx, column=12)
                    c_pr.alignment = Alignment(horizontal="center", vertical="center")
                    c_pr.border = cell_border

                ws_bm.freeze_panes = "A2"
                ws_bm.auto_filter.ref = f"A1:L{len(rows) + 1}"

                col_widths = {
                    "A": 12, "B": 15, "C": 12, "D": 16, "E": 15,
                    "F": 15, "G": 16, "H": 16, "I": 18, "J": 14,
                    "K": 13, "L": 11
                }
                for col_letter, width in col_widths.items():
                    ws_bm.column_dimensions[col_letter].width = width

        else:
            # Single benchmark view
            ws1.title = "Correlation Analysis"

            headers = [
                "Ticker",
                "Correlation",
                "Beta",
                "Expected Move",
                "Actual Move",
                "Difference",
                "Normalized Diff",
                "Daily Volatility",
                "Residual Volatility",
                "Observations",
                "Benchmark",
                "Period"
            ]
            ws1.append(headers)
            ws1.row_dimensions[1].height = 26

            for col_num in range(1, len(headers) + 1):
                cell = ws1.cell(row=1, column=col_num)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_alignment
                cell.border = cell_border

            for row_idx, r in enumerate(rows, start=2):
                ticker = str(r.get("ticker", "")).strip().upper()
                corr = clean_value(r.get("correlation"))
                beta_val = clean_value(r.get("beta"))
                exp_move = clean_value(r.get("expected_move"))
                act_move = clean_value(r.get("actual_move"))
                diff = clean_value(r.get("difference"))
                norm = clean_value(r.get("normalized"))
                vol = clean_value(r.get("volatility"))
                res_vol = clean_value(r.get("residual_volatility"))
                obs = clean_value(r.get("observations"))

                row_data = [
                    ticker,
                    corr,
                    beta_val,
                    exp_move,
                    act_move,
                    diff,
                    norm,
                    vol,
                    res_vol,
                    obs,
                    benchmark,
                    period
                ]
                ws1.append(row_data)
                ws1.row_dimensions[row_idx].height = 20

                c_ticker = ws1.cell(row=row_idx, column=1)
                c_ticker.alignment = Alignment(horizontal="center", vertical="center")
                c_ticker.font = Font(name="Calibri", size=10, bold=True)
                c_ticker.border = cell_border

                c_corr = ws1.cell(row=row_idx, column=2)
                c_corr.alignment = Alignment(horizontal="right", vertical="center")
                c_corr.border = cell_border
                format_corr_cell(c_corr, corr)

                c_beta = ws1.cell(row=row_idx, column=3)
                c_beta.alignment = Alignment(horizontal="right", vertical="center")
                c_beta.border = cell_border
                if beta_val is not None:
                    c_beta.number_format = "0.000"

                c_exp = ws1.cell(row=row_idx, column=4)
                c_exp.alignment = Alignment(horizontal="right", vertical="center")
                c_exp.border = cell_border
                if exp_move is not None:
                    c_exp.number_format = "0.00%"

                c_act = ws1.cell(row=row_idx, column=5)
                c_act.alignment = Alignment(horizontal="right", vertical="center")
                c_act.border = cell_border
                if act_move is not None:
                    c_act.number_format = "0.00%"

                c_diff = ws1.cell(row=row_idx, column=6)
                c_diff.alignment = Alignment(horizontal="right", vertical="center")
                c_diff.border = cell_border
                if diff is not None:
                    c_diff.number_format = "0.00%"
                    if diff > 0:
                        c_diff.font = Font(name="Calibri", size=10, color="16A34A")
                    elif diff < 0:
                        c_diff.font = Font(name="Calibri", size=10, color="DC2626")

                c_norm = ws1.cell(row=row_idx, column=7)
                c_norm.alignment = Alignment(horizontal="right", vertical="center")
                c_norm.border = cell_border
                if norm is not None:
                    c_norm.number_format = "0.00"

                c_vol = ws1.cell(row=row_idx, column=8)
                c_vol.alignment = Alignment(horizontal="right", vertical="center")
                c_vol.border = cell_border
                if vol is not None:
                    c_vol.number_format = "0.00%"

                c_res = ws1.cell(row=row_idx, column=9)
                c_res.alignment = Alignment(horizontal="right", vertical="center")
                c_res.border = cell_border
                if res_vol is not None:
                    c_res.number_format = "0.00%"

                c_obs = ws1.cell(row=row_idx, column=10)
                c_obs.alignment = Alignment(horizontal="center", vertical="center")
                c_obs.border = cell_border
                if obs is not None:
                    c_obs.number_format = "#,##0"

                c_bm = ws1.cell(row=row_idx, column=11)
                c_bm.alignment = Alignment(horizontal="center", vertical="center")
                c_bm.border = cell_border

                c_per = ws1.cell(row=row_idx, column=12)
                c_per.alignment = Alignment(horizontal="center", vertical="center")
                c_per.border = cell_border

            ws1.freeze_panes = "A2"
            ws1.auto_filter.ref = f"A1:L{len(rows) + 1}"

            col_widths = {
                "A": 12, "B": 15, "C": 12, "D": 16, "E": 15,
                "F": 15, "G": 16, "H": 16, "I": 18, "J": 14,
                "K": 13, "L": 11
            }
            for col_letter, width in col_widths.items():
                ws1.column_dimensions[col_letter].width = width

        # ========================================================
        # SHEET: Summary & Export Metadata
        # ========================================================
        ws_summary = wb.create_sheet(title="Summary & Metadata")
        ws_summary.views.sheetView[0].showGridLines = True

        ws_summary.append(["Stock Correlation & Portfolio Analytics Summary"])
        ws_summary.cell(row=1, column=1).font = Font(name="Calibri", size=14, bold=True, color="1E293B")
        ws_summary.append([])

        metadata_rows = [
            ["Market Benchmark(s)", ", ".join(benchmarks)],
            ["Analysis Period", period],
            ["Filter Criteria", filter_label],
            ["Exported Assets Count", len(rows)],
            ["Export Timestamp (UTC)", datetime.datetime.utcnow().strftime("%Y-%m-%d %H:%M:%S")]
        ]

        # Calculate average correlations and betas per benchmark
        for bm in benchmarks:
            bm_corrs = []
            bm_betas = []
            for r in rows:
                c = clean_value(r.get(f"correlation_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("correlation") or r.get("correlation"))
                b = clean_value(r.get(f"beta_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("beta") or r.get("beta"))
                if c is not None:
                    bm_corrs.append(c)
                if b is not None:
                    bm_betas.append(b)

            avg_c = float(np.mean(bm_corrs)) if bm_corrs else None
            avg_b = float(np.mean(bm_betas)) if bm_betas else None
            metadata_rows.append([f"Avg Correlation vs {bm}", f"{avg_c:.3f}" if avg_c is not None else "—"])
            metadata_rows.append([f"Avg Beta vs {bm}", f"{avg_b:.3f}" if avg_b is not None else "—"])

        meta_label_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        meta_label_font = Font(name="Calibri", size=10, bold=True, color="334155")

        for row_data in metadata_rows:
            ws_summary.append(row_data)
            curr_row = ws_summary.max_row
            cell_lbl = ws_summary.cell(row=curr_row, column=1)
            cell_val = ws_summary.cell(row=curr_row, column=2)

            cell_lbl.fill = meta_label_fill
            cell_lbl.font = meta_label_font
            cell_lbl.border = cell_border

            cell_val.font = Font(name="Calibri", size=10)
            cell_val.border = cell_border

        ws_summary.column_dimensions["A"].width = 28
        ws_summary.column_dimensions["B"].width = 32

        # Save workbook to memory
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        clean_bm = "_".join("".join(c for c in bm if c.isalnum()) for bm in benchmarks[:3]) or "benchmark"
        clean_period = "".join(c for c in period if c.isalnum()) or "1y"
        filename = f"correlation_{clean_bm}_{clean_period}.xlsx"

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )

    except Exception as e:
        print("Excel export error:", e)
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# ============================================================
# DEDICATED CUSTOM MOVE & RELATIVE PERFORMANCE TRACKER
# ============================================================

@app.route("/api/calculate-custom-tracker", methods=["POST"])
def calculate_custom_tracker():
    try:
        data = request.get_json() or {}
        items = data.get("items", []) # list of dicts: [{"ticker": "COHR", "actual_move": 0.04}, ...]
        reference = str(data.get("reference") or data.get("benchmark") or "SPY").strip().upper()
        reference_move = data.get("reference_move", 0.01)
        period = str(data.get("period", "1y"))

        try:
            reference_move = float(reference_move)
        except (TypeError, ValueError):
            reference_move = 0.01

        if not items:
            return jsonify({"rows": [], "reference": reference, "reference_move": reference_move, "period": period})

        # Clean items & moves
        parsed_items = {}
        for item in items:
            if isinstance(item, dict):
                sym = str(item.get("ticker", "")).strip().upper()
                mv = item.get("actual_move")
            else:
                sym = str(item).strip().upper()
                mv = None
            if sym and sym != reference:
                try:
                    mv_val = float(mv) if mv is not None else None
                except (TypeError, ValueError):
                    mv_val = None
                parsed_items[sym] = mv_val

        tickers_list = list(parsed_items.keys())
        if not tickers_list:
            return jsonify({"rows": [], "reference": reference, "reference_move": reference_move, "period": period})

        all_symbols = [reference] + tickers_list
        df = get_stock_data(*all_symbols)

        if df is None or df.empty:
            return jsonify({"error": "Данные для указанных тикеров не найдены в базе данных."}), 404

        available_tickers = set(
            df["ticker"].astype(str).str.strip().str.upper().unique()
        )

        if reference not in available_tickers:
            return jsonify({"error": f"Базовый актив '{reference}' не найден в базе данных."}), 404

        valid_tickers = [t for t in tickers_list if t in available_tickers]
        if not valid_tickers:
            return jsonify({"error": "Ни один из указанных тикеров не найден в базе данных."}), 404

        analysis = calculate_groups_analysis(
            df=df,
            groups={"CustomTracker": valid_tickers},
            benchmarks=[reference],
            period=period,
            benchmark_move=reference_move,
            asset_moves={t: parsed_items[t] for t in valid_tickers if parsed_items[t] is not None}
        )

        ignored_tickers = analysis.get("ignored_tickers", {})
        partial_tickers = analysis.get("partial_tickers", {})
        bm_data = analysis.get("benchmarks_data", {}).get(reference, {})
        expected_range = bm_data.get("expected_range", {})
        correlation_matrix = analysis.get("correlation", {})
        analyzed_tickers = analysis.get("tickers", [])

        rows = []
        for ticker in analyzed_tickers:
            exp_info = expected_range.get(ticker, {})
            corr_val = correlation_matrix.get(ticker, {}).get(reference)
            act_mv = parsed_items.get(ticker)
            if act_mv is None and exp_info.get("actual_move") is not None:
                act_mv = exp_info.get("actual_move")

            rows.append({
                "ticker": ticker,
                "correlation": clean_value(corr_val),
                "beta": clean_value(exp_info.get("beta")),
                "expected_move": clean_value(exp_info.get("expected_stock_move")),
                "actual_move": clean_value(act_mv),
                "difference": clean_value(exp_info.get("difference")),
                "normalized": clean_value(exp_info.get("normalized_difference")),
                "volatility": clean_value(exp_info.get("volatility")),
                "residual_volatility": clean_value(exp_info.get("residual_volatility")),
                "observations": clean_value(exp_info.get("observations")),
                "reference": reference,
                "period": period,
                "is_partial_history": bool(exp_info.get("is_partial_history", False)),
                "history_period_ru": exp_info.get("history_period_ru"),
                "history_text": exp_info.get("history_text"),
                "history_start": exp_info.get("history_start"),
                "history_end": exp_info.get("history_end"),
                "history_days": clean_value(exp_info.get("history_days")),
                "history_observations": clean_value(exp_info.get("history_observations")),
                "selected_period": exp_info.get("selected_period", period),
                "selected_period_label": exp_info.get("selected_period_label"),
            })

        ignored_list = list(ignored_tickers.keys()) if isinstance(ignored_tickers, dict) else list(ignored_tickers)
        partial_list = list(partial_tickers.keys()) if isinstance(partial_tickers, dict) else list(partial_tickers)

        return jsonify({
            "rows": rows,
            "reference": reference,
            "reference_move": reference_move,
            "period": period,
            "missing": [t for t in tickers_list if t not in available_tickers],
            "ignored": ignored_list,
            "ignored_tickers": clean_dict(ignored_tickers),
            "partial": partial_list,
            "partial_tickers": clean_dict(partial_tickers),
        })

    except Exception as e:
        print("Custom tracker calculation error:", e)
        import traceback
        traceback.print_exc()
        return jsonify({"error": str(e)}), 500


# ============================================================
# DATABASE STATUS
# ============================================================

@app.route("/api/status")
@app.route("/api/db-status")
def db_status():
    try:
        return jsonify(get_database_status())
    except Exception as e:
        return jsonify({
            "online": False,
            "source": "error",
            "message": str(e)
        }), 500


# ============================================================
# TICKERS LIST
# ============================================================

@app.route("/api/tickers")
def tickers():
    try:
        return jsonify(get_available_tickers())
    except Exception as e:
        return jsonify({"error": str(e)}), 500


def compute_stock_volatility(ticker: str):
    try:
        df_daily = get_stock_data(ticker)
        if df_daily.empty:
            poly_bars = polygon_client.fetch_daily_bars(ticker, start_date="2023-01-01")
            if poly_bars:
                df_daily = pd.DataFrame(poly_bars)
        if not df_daily.empty and len(df_daily) >= 15:
            df_sorted = df_daily.sort_values("date").tail(252)
            rets = df_sorted["close"].pct_change().dropna()
            if len(rets) >= 10:
                ann_vol = float(rets.std() * np.sqrt(252) * 100.0)
                return round(ann_vol, 1)
    except Exception as e:
        print(f"Error computing volatility for {ticker}:", e)
    return None


@app.route("/api/ticker/<ticker>/profile")
def flask_ticker_profile(ticker):
    ticker = str(ticker).strip().upper()
    cached = get_company_profile_db(ticker)
    vol = compute_stock_volatility(ticker)
    if cached:
        cached["volatility"] = vol or cached.get("volatility")
        cached["description"] = cached.get("description_ru")
        return jsonify(cached)

    profile = fmp_client.get_company_profile(ticker)
    peers = fmp_client.get_stock_peers(ticker)
    profile["peers"] = [p["symbol"] for p in peers if p.get("symbol")]
    profile["volatility"] = vol
    profile["description"] = profile.get("description") or profile.get("description_ru")

    save_company_profile_db(ticker, profile)
    return jsonify(profile)


@app.route("/api/ticker/<ticker>/chart/daily")
def flask_daily_chart(ticker):
    ticker = str(ticker).strip().upper()

    # 1. Fetch daily bars
    df = get_stock_data(ticker)
    if df.empty:
        poly_bars = polygon_client.fetch_daily_bars(ticker, start_date="2021-01-01")
        if poly_bars:
            df = pd.DataFrame(poly_bars)
            df["date"] = pd.to_datetime(df["date"])

    candles = []
    if not df.empty:
        df = df.sort_values("date").reset_index(drop=True)
        for _, r in df.iterrows():
            d_str = pd.to_datetime(r["date"]).strftime("%Y-%m-%d")
            candles.append({
                "time": d_str,
                "open": float(r["open"]),
                "high": float(r["high"]),
                "low": float(r["low"]),
                "close": float(r["close"]),
                "volume": float(r.get("volume", 0.0))
            })

    # 2. Markers
    markers = []
    marker_details = {}

    candle_dates = {c["time"] for c in candles}
    sorted_candle_dates = sorted(list(candle_dates))

    def find_nearest_candle_date(target_date: str):
        if target_date in candle_dates:
            return target_date
        if not sorted_candle_dates or target_date < sorted_candle_dates[0] or target_date > sorted_candle_dates[-1]:
            return None
        import bisect
        idx = bisect.bisect_left(sorted_candle_dates, target_date)
        if idx < len(sorted_candle_dates):
            cand = sorted_candle_dates[idx]
            try:
                d_target = datetime.date.fromisoformat(target_date)
                d_cand = datetime.date.fromisoformat(cand)
                if abs((d_cand - d_target).days) <= 4:
                    return cand
            except Exception:
                pass
        return None

    seen_rep_dates = set()
    extra_map = get_ticker_earnings_extra_map(ticker)
    fmp_reports = fmp_client.get_earnings_history(ticker, limit=30)
    for fr in fmp_reports:
        raw_d = fr.get("date")
        if not raw_d:
            continue
        c_d = find_nearest_candle_date(raw_d)
        if not c_d or c_d in seen_rep_dates:
            continue
        seen_rep_dates.add(c_d)
        extra_info = find_matching_earnings_extra(extra_map, raw_d)
        markers.append({
            "time": c_d,
            "position": "belowBar",
            "color": "#10b981",
            "shape": "square",
            "text": "E",
            "id": f"e_{c_d}"
        })
        marker_details[f"e_{c_d}"] = {
            "type": "Earnings",
            "date": c_d,
            "session": str(fr.get("time", "AMC")).upper(),
            "eps_actual": fr.get("eps_actual") if fr.get("eps_actual") is not None else (extra_info.get("eps_actual") if extra_info else None),
            "eps_estimate": fr.get("eps_estimate") if fr.get("eps_estimate") is not None else (extra_info.get("eps_estimate") if extra_info else None),
            "eps_surprise_pct": fr.get("eps_surprise_pct") if fr.get("eps_surprise_pct") is not None else (extra_info.get("eps_surprise_pct") if extra_info else None),
            "revenue_actual": fr.get("revenue_actual") if fr.get("revenue_actual") is not None else (extra_info.get("revenue_actual") if extra_info else None),
            "revenue_estimate": fr.get("revenue_estimate") if fr.get("revenue_estimate") is not None else (extra_info.get("revenue_estimate") if extra_info else None)
        }

    db_reports = get_ticker_reports(ticker)
    if not db_reports.empty:
        for _, r in db_reports.iterrows():
            raw_d = str(r["date"])
            c_d = find_nearest_candle_date(raw_d)
            if not c_d and r.get("move_date"):
                c_d = find_nearest_candle_date(str(r["move_date"]))
            if not c_d or c_d in seen_rep_dates:
                continue
            seen_rep_dates.add(c_d)
            time_str = str(r.get("time", "")).lower()
            session_name = "AMC" if "16:" in time_str or "amc" in time_str or "after" in time_str else "BMO"
            quarter_name = f"{r.get('fiscal_period', '')} {r.get('fiscal_year', '')}".strip()
            extra_info = find_matching_earnings_extra(extra_map, raw_d)
            if not extra_info and r.get("move_date"):
                extra_info = find_matching_earnings_extra(extra_map, str(r["move_date"]))

            markers.append({
                "time": c_d,
                "position": "belowBar",
                "color": "#10b981",
                "shape": "square",
                "text": "E",
                "id": f"e_{c_d}"
            })
            marker_details[f"e_{c_d}"] = {
                "type": "Earnings",
                "date": c_d,
                "session": (extra_info.get("time") if extra_info else None) or session_name,
                "quarter": quarter_name,
                "eps_actual": extra_info.get("eps_actual") if extra_info else None,
                "eps_estimate": extra_info.get("eps_estimate") if extra_info else None,
                "eps_surprise_pct": extra_info.get("eps_surprise_pct") if extra_info else None,
                "revenue_actual": extra_info.get("revenue_actual") if extra_info else None,
                "revenue_estimate": extra_info.get("revenue_estimate") if extra_info else None
            }

    # Dividends
    seen_div_dates = set()
    db_divs = get_ticker_dividends_db(ticker)
    if not db_divs:
        db_divs = fmp_client.get_dividends_history(ticker, limit=30)
        if not db_divs:
            db_divs = polygon_client.fetch_dividends(ticker)
        if db_divs:
            save_ticker_dividends_db(ticker, db_divs)

    for div in db_divs:
        raw_ex = div.get("ex_date") or div.get("date")
        if not raw_ex:
            continue
        c_ex = find_nearest_candle_date(raw_ex)
        if not c_ex or c_ex in seen_div_dates:
            continue
        seen_div_dates.add(c_ex)
        markers.append({
            "time": c_ex,
            "position": "belowBar",
            "color": "#d97706",
            "shape": "circle",
            "text": "D",
            "id": f"d_{c_ex}"
        })
        marker_details[f"d_{c_ex}"] = {
            "type": "Dividend",
            "ex_date": c_ex,
            "amount": div.get("amount") or div.get("dividend"),
            "pay_date": div.get("pay_date") or div.get("paymentDate"),
            "yield": div.get("yield"),
            "frequency": div.get("frequency", "Quarterly")
        }

    markers.sort(key=lambda x: x["time"])
    return jsonify({
        "ticker": ticker,
        "candles": candles,
        "markers": markers,
        "marker_details": marker_details
    })


@app.route("/api/ticker/<ticker>/chart/intraday")
def flask_intraday_chart(ticker):
    ticker = str(ticker).strip().upper()
    date = request.args.get("date")
    multiplier = int(request.args.get("multiplier", 5))
    days_back = int(request.args.get("days_back", 30))

    today_dt = datetime.date.today()
    if date:
        try:
            req_dt = datetime.date.fromisoformat(date)
        except Exception:
            req_dt = today_dt
    else:
        req_dt = today_dt

    start_date = (req_dt - datetime.timedelta(days=days_back)).isoformat()
    end_date = min(today_dt, req_dt + datetime.timedelta(days=10)).isoformat()

    df = get_intraday_chart_data(ticker, start_date=start_date, end_date=end_date)
    candles = []

    if not df.empty:
        df = df.sort_values(["date", "time"])
        for _, r in df.iterrows():
            dt_str = f"{r['date']} {r['time']}"
            try:
                dt_obj = datetime.datetime.strptime(dt_str, "%Y-%m-%d %H:%M:%S")
                ts = int(dt_obj.replace(tzinfo=datetime.timezone.utc).timestamp())
            except Exception:
                ts = None
            candles.append({
                "time": dt_str,
                "timestamp": ts,
                "date": str(r["date"]),
                "time_str": r["time_str"],
                "open": float(r["open"]),
                "high": float(r["high"]),
                "low": float(r["low"]),
                "close": float(r["close"]),
                "volume": float(r["volume"]),
                "session_type": r.get("session_type", "regular")
            })
    else:
        poly_bars = polygon_client.fetch_intraday_bars(
            ticker,
            start_date=start_date,
            end_date=end_date,
            multiplier=multiplier,
            timespan="minute"
        )
        for b in poly_bars:
            candles.append({
                "time": b["datetime_et"],
                "timestamp": b.get("timestamp"),
                "date": b.get("date"),
                "time_str": b["time_str"],
                "open": b["open"],
                "high": b["high"],
                "low": b["low"],
                "close": b["close"],
                "volume": b["volume"],
                "session_type": b["session_type"]
            })

    return jsonify({
        "ticker": ticker,
        "date": req_dt.isoformat(),
        "start_date": start_date,
        "end_date": end_date,
        "multiplier": multiplier,
        "candles": candles,
        "count": len(candles)
    })


@app.route("/api/ticker/<ticker>/earnings-moves")
def flask_earnings_moves(ticker):
    ticker = str(ticker).strip().upper()
    window = request.args.get("window", "04:00_09:30")
    limit = int(request.args.get("limit", 8))

    matrix = earnings_reaction_client.get_earnings_reaction_matrix(
        ticker,
        limit=limit,
        premarket_window=window
    )
    return jsonify({
        "ticker": ticker,
        "window": window,
        "count": len(matrix),
        "matrix": matrix
    })


@app.route("/api/ticker/<ticker>/earnings-moves/export")
def flask_export_earnings_moves(ticker):
    ticker = str(ticker).strip().upper()
    r_minus_1_m = request.args.get("r_minus_1_metric", "postmarket")
    r_0_m = request.args.get("r_0_metric", "close_to_close")
    r_plus_1_m = request.args.get("r_plus_1_metric", "close_to_close")
    r_plus_2_m = request.args.get("r_plus_2_metric", "close_to_close")

    matrix = earnings_reaction_client.get_earnings_reaction_matrix(ticker, limit=8)
    if not matrix:
        return jsonify({"error": f"Нет данных об отчетах для тикера {ticker}"}), 404

    metric_labels = {
        "close_to_close": "Close to Close %",
        "premarket": "Premarket %",
        "intraday": "Intraday %",
        "postmarket": "Postmarket %"
    }

    selected_cols = [
        ("R-1", r_minus_1_m, f"R-1 ({metric_labels.get(r_minus_1_m, r_minus_1_m)})"),
        ("R", r_0_m, f"R ({metric_labels.get(r_0_m, r_0_m)})"),
        ("R+1", r_plus_1_m, f"R+1 ({metric_labels.get(r_plus_1_m, r_plus_1_m)})"),
        ("R+2", r_plus_2_m, f"R+2 ({metric_labels.get(r_plus_2_m, r_plus_2_m)})"),
    ]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{ticker} Earnings Moves"
    ws.views.sheetView[0].showGridLines = True

    # Title Banner
    ws.merge_cells("A1:J1")
    title_cell = ws["A1"]
    title_cell.value = f"Матрица движений на отчетности: {ticker} (Последние 8 кварталов)"
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
    title_cell.fill = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36

    headers = [
        "Квартал", "Дата отчета", "Сессия", "EPS Факт / Ожид.", "Сюрприз EPS", "Выручка Факт / Ожид.",
        selected_cols[0][2], selected_cols[1][2], selected_cols[2][2], selected_cols[3][2]
    ]

    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2D3748", end_color="2D3748", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    ws.row_dimensions[2].height = 26
    for col_idx, text in enumerate(headers, start=1):
        cell = ws.cell(row=2, column=col_idx, value=text)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    green_fill = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    green_font = Font(name="Calibri", size=10, color="166534", bold=True)
    red_fill = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
    red_font = Font(name="Calibri", size=10, color="991B1B", bold=True)
    standard_font = Font(name="Calibri", size=10)

    day_stats = {"R-1": [], "R": [], "R+1": [], "R+2": []}

    for row_idx, r in enumerate(matrix, start=3):
        ws.row_dimensions[row_idx].height = 20
        eps_str = f"${r['eps_actual']:.2f} / ${r['eps_estimate']:.2f}" if r.get("eps_actual") is not None and r.get("eps_estimate") is not None else "—"
        rev_act_f = f"${r['revenue_actual']/1e9:.2f}B" if r.get("revenue_actual") and r.get("revenue_actual") > 1e9 else (f"${r['revenue_actual']/1e6:.1f}M" if r.get("revenue_actual") else "—")
        rev_est_f = f"${r['revenue_estimate']/1e9:.2f}B" if r.get("revenue_estimate") and r.get("revenue_estimate") > 1e9 else (f"${r['revenue_estimate']/1e6:.1f}M" if r.get("revenue_estimate") else "—")
        rev_str = f"{rev_act_f} / {rev_est_f}"

        row_vals = [r.get("quarter", ""), r.get("date", ""), r.get("session", "AMC"), eps_str]
        ws.cell(row=row_idx, column=1, value=row_vals[0]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=row_vals[1]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=3, value=row_vals[2]).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=row_vals[3]).alignment = Alignment(horizontal="center")

        surp_cell = ws.cell(row=row_idx, column=5)
        if r.get("eps_surprise_pct") is not None:
            val = float(r["eps_surprise_pct"])
            surp_cell.value = f"{'+' if val >= 0 else ''}{val:.2f}%"
            surp_cell.fill = green_fill if val >= 0 else red_fill
            surp_cell.font = green_font if val >= 0 else red_font
        else:
            surp_cell.value = "—"
        surp_cell.alignment = Alignment(horizontal="center")

        ws.cell(row=row_idx, column=6, value=rev_str).alignment = Alignment(horizontal="center")

        for c_idx, (day_k, met_k, _) in enumerate(selected_cols, start=7):
            cell = ws.cell(row=row_idx, column=c_idx)
            val = r.get("days", {}).get(day_k, {}).get(met_k)
            if val is not None and not np.isnan(val):
                val_num = float(val)
                day_stats[day_k].append(val_num)
                cell.value = f"{'+' if val_num >= 0 else ''}{val_num:.2f}%"
                cell.fill = green_fill if val_num >= 0 else red_fill
                cell.font = green_font if val_num >= 0 else red_font
            else:
                cell.value = "—"
            cell.alignment = Alignment(horizontal="center")

        for col_i in range(1, 11):
            c = ws.cell(row=row_idx, column=col_i)
            c.border = thin_border
            if not c.font.bold:
                c.font = standard_font

    # Footer Statistics
    stat_rows = [
        ("MEAN (СРЕДНЕЕ)", lambda vals: np.mean(vals) if len(vals) else None),
        ("MEDIAN (МЕДИАНА)", lambda vals: np.median(vals) if len(vals) else None),
        ("MAX (МАКСИМУМ)", lambda vals: np.max(vals) if len(vals) else None),
        ("MIN (МИНИМУМ)", lambda vals: np.min(vals) if len(vals) else None),
        ("WIN RATE (% ПОЛОЖИТЕЛЬНЫХ)", lambda vals: (sum(1 for v in vals if v > 0) / len(vals) * 100) if len(vals) else None)
    ]

    stat_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    stat_font_lbl = Font(name="Calibri", size=10, bold=True, color="334155")
    start_stat_row = 3 + len(matrix)

    for s_idx, (label, func) in enumerate(stat_rows, start=start_stat_row):
        ws.row_dimensions[s_idx].height = 20
        ws.merge_cells(start_row=s_idx, start_column=1, end_row=s_idx, end_column=6)
        lbl_cell = ws.cell(row=s_idx, column=1, value=label)
        lbl_cell.font = stat_font_lbl
        lbl_cell.fill = stat_fill
        lbl_cell.alignment = Alignment(horizontal="right", vertical="center")

        for col_i in range(1, 7):
            ws.cell(row=s_idx, column=col_i).border = thin_border
            ws.cell(row=s_idx, column=col_i).fill = stat_fill

        for c_idx, (day_k, _, _) in enumerate(selected_cols, start=7):
            cell = ws.cell(row=s_idx, column=c_idx)
            vals = day_stats.get(day_k, [])
            calc = func(vals)
            if calc is not None:
                if "WIN RATE" in label:
                    cell.value = f"{calc:.0f}%"
                    cell.font = green_font
                else:
                    cell.value = f"{'+' if calc >= 0 else ''}{calc:.2f}%"
                    cell.font = green_font if calc >= 0 else red_font
                    cell.fill = green_fill if calc >= 0 else red_fill
            else:
                cell.value = "—"
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border

    # Auto-fit columns
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"{ticker}_Earnings_Moves_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename
    )


@app.route("/api/ticker/<ticker>/correlations")
def flask_ticker_correlations(ticker):
    ticker = str(ticker).strip().upper()
    period = request.args.get("period", "1y")
    peers_param = request.args.get("peers")
    bm_param = request.args.get("benchmarks")

    custom_peers = [p.strip().upper() for p in peers_param.split(",") if p.strip()] if peers_param is not None else None
    custom_bm = [b.strip().upper() for b in bm_param.split(",") if b.strip()] if bm_param is not None else None

    result = quant_client.calculate_ticker_correlations(
        ticker=ticker,
        period=period,
        peers=custom_peers,
        benchmarks=custom_bm
    )
    return jsonify(result)


@app.route("/api/ticker/<ticker>/dividends")
def flask_dividends(ticker):
    ticker = str(ticker).strip().upper()
    db_divs = get_ticker_dividends_db(ticker)
    if not db_divs:
        db_divs = fmp_client.get_dividends_history(ticker, limit=24)
        if not db_divs:
            db_divs = polygon_client.fetch_dividends(ticker, limit=24)
        if db_divs:
            save_ticker_dividends_db(ticker, db_divs)

    return jsonify({
        "ticker": ticker,
        "count": len(db_divs),
        "dividends": db_divs
    })


@app.route("/api/ticker/<ticker>/sync", methods=["POST"])
def flask_sync_ticker(ticker):
    ticker = str(ticker).strip().upper()
    prof = fmp_client.get_company_profile(ticker)
    peers = fmp_client.get_stock_peers(ticker)
    prof["peers"] = [p["symbol"] for p in peers if p.get("symbol")]
    save_company_profile_db(ticker, prof)

    divs = fmp_client.get_dividends_history(ticker, limit=24)
    if not divs:
        divs = polygon_client.fetch_dividends(ticker, limit=24)
    if divs:
        save_ticker_dividends_db(ticker, divs)

    return jsonify({
        "ticker": ticker,
        "profile": prof.get("companyName"),
        "peers": len(prof["peers"]),
        "dividends": len(divs) if divs else 0,
        "status": "success"
    })


if __name__ == "__main__":
    app.run(
        host="127.0.0.1",
        port=8000,
        debug=True
    )