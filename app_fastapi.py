import os
import datetime
from typing import Dict, List, Any, Optional
import pandas as pd
import numpy as np
from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, JSONResponse, FileResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

import io
import re
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from config import BASE_DIR, BENCHMARKS
from database import (
    get_available_tickers,
    get_database_status,
    get_stock_data,
    get_intraday_chart_data,
    get_ticker_reports,
    get_ticker_dividends_db,
    save_ticker_dividends_db,
    get_company_profile_db,
    save_company_profile_db
)
from calculations import (
    calculate_all_periods,
    calculate_groups_analysis
)
from services.polygon_service import polygon_client
from services.fmp_service import fmp_client
from services.earnings_reaction_service import earnings_reaction_client
from services.quant_service import quant_client
from services.earnings_extra_service import get_ticker_earnings_extra_map, find_matching_earnings_extra

app = FastAPI(
    title="Helion Financial Analysis Platform",
    description="Quantitative ticker analytics, dual TradingView charts, earnings moves, dividends, and correlation matrix.",
    version="2.0.0"
)

# CORS
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Static files directory
STATIC_DIR = BASE_DIR / "static"
os.makedirs(STATIC_DIR, exist_ok=True)
app.mount("/static", StaticFiles(directory=str(STATIC_DIR)), name="static")


# ============================================================
# ROOT PAGE
# ============================================================

@app.get("/", response_class=HTMLResponse)
async def get_index():
    index_file = STATIC_DIR / "index.html"
    if index_file.exists():
        return FileResponse(str(index_file))
    return HTMLResponse("<h1>Helion Quant Analysis Platform API Running</h1><p>Visit /docs for Swagger API</p>")


# ============================================================
# SYSTEM STATUS & TICKERS
# ============================================================

@app.get("/api/status")
async def api_status():
    db_stat = get_database_status()
    return {
        "status": "online",
        "database": db_stat,
        "timestamp": datetime.datetime.now().isoformat(),
        "benchmarks": BENCHMARKS
    }


@app.get("/api/tickers")
async def api_tickers(q: Optional[str] = None):
    all_tickers = get_available_tickers()
    if q:
        query = q.strip().upper()
        filtered = [t for t in all_tickers if query in t]
        return {"tickers": filtered[:50], "total": len(filtered)}
    return {"tickers": all_tickers, "total": len(all_tickers)}


# ============================================================
# COMPANY PROFILE & PEERS
# ============================================================

def compute_stock_volatility(ticker: str) -> Optional[float]:
    try:
        df_daily = get_stock_data(ticker)
        if df_daily.empty:
            poly_bars = polygon_client.fetch_daily_bars(ticker, start_date="2023-01-01")
            if poly_bars:
                df_daily = pd.DataFrame(poly_bars)
        if not df_daily.empty and len(df_daily) >= 15:
            df_sorted = df_daily.sort_values("date").tail(252)
            rets = df_sorted["close"].pct_change().dropna()
            if len(rets) >= 10:
                ann_vol = float(rets.std() * np.sqrt(252) * 100.0)
                return round(ann_vol, 1)
    except Exception as e:
        print(f"Error computing volatility for {ticker}:", e)
    return None


@app.get("/api/ticker/{ticker}/profile")
async def api_company_profile(ticker: str):
    ticker = ticker.strip().upper()
    cached = get_company_profile_db(ticker)
    vol = compute_stock_volatility(ticker)
    
    peers_list = fmp_client.get_stock_peers(ticker)
    business_peers = [
        {
            "symbol": p["symbol"].strip().upper(),
            "companyName": p.get("companyName", p["symbol"]),
            "price": p.get("price")
        }
        for p in peers_list if p.get("symbol") and p.get("symbol").strip().upper() != ticker
    ]
    peer_symbols = [p["symbol"] for p in business_peers]

    if cached and cached.get("description_ru"):
        raw_db_peers = cached.get("peers", [])
        if isinstance(raw_db_peers, str):
            raw_db_peers = [p.strip() for p in raw_db_peers.split(",") if p.strip()]
        cached["peers"] = peer_symbols or raw_db_peers
        cached["business_peers"] = business_peers
        cached["volatility"] = vol or cached.get("volatility")
        cached["description"] = cached.get("description_ru")
        return cached

    profile = fmp_client.get_company_profile(ticker)
    profile["peers"] = peer_symbols
    profile["business_peers"] = business_peers
    profile["volatility"] = vol
    profile["description"] = profile.get("description") or profile.get("description_ru")

    save_company_profile_db(ticker, profile)
    return profile


# ============================================================
# DUAL CHARTS DATA
# ============================================================

@app.get("/api/ticker/{ticker}/chart/daily")
async def api_daily_chart(ticker: str):
    """
    Returns daily OHLCV candlestick series + Earnings 'E' and Dividends 'D' event markers.
    """
    ticker = ticker.strip().upper()

    # 1. Fetch daily candles from DB or Polygon
    df = get_stock_data(ticker)
    candles = []

    if not df.empty:
        df["date_str"] = df["date"].dt.strftime("%Y-%m-%d")
        df = df.sort_values("date_str")
        for _, r in df.iterrows():
            candles.append({
                "time": r["date_str"],
                "open": float(r["open"]),
                "high": float(r["high"]),
                "low": float(r["low"]),
                "close": float(r["close"]),
                "volume": float(r["volume"]) if pd.notnull(r["volume"]) else 0.0
            })
    else:
        poly_bars = polygon_client.fetch_daily_bars(ticker, start_date="2020-01-01")
        for b in poly_bars:
            candles.append({
                "time": b["date"],
                "open": b["open"],
                "high": b["high"],
                "low": b["low"],
                "close": b["close"],
                "volume": b["volume"]
            })

    # 2. Build Event Markers
    markers = []
    marker_details = {}

    candle_dates = {c["time"] for c in candles}
    sorted_candle_dates = sorted(list(candle_dates))

    def find_nearest_candle_date(target_date: str) -> Optional[str]:
        if target_date in candle_dates:
            return target_date
        if not sorted_candle_dates or target_date < sorted_candle_dates[0] or target_date > sorted_candle_dates[-1]:
            return None
        import bisect
        idx = bisect.bisect_left(sorted_candle_dates, target_date)
        if idx < len(sorted_candle_dates):
            cand = sorted_candle_dates[idx]
            try:
                d_target = datetime.date.fromisoformat(target_date)
                d_cand = datetime.date.fromisoformat(cand)
                if abs((d_cand - d_target).days) <= 4:
                    return cand
            except Exception:
                pass
        return None

    # A. Earnings Markers ('E')
    seen_earnings_dates = set()
    extra_map = get_ticker_earnings_extra_map(ticker)
    earnings_list = fmp_client.get_earnings_history(ticker, limit=30)
    for earn in earnings_list:
        raw_date = earn.get("date")
        if not raw_date:
            continue
        c_date = find_nearest_candle_date(raw_date)
        if not c_date or c_date in seen_earnings_dates:
            continue
        seen_earnings_dates.add(c_date)
        m_id = f"e_{c_date}"
        extra_info = find_matching_earnings_extra(extra_map, raw_date)
        markers.append({
            "time": c_date,
            "position": "belowBar",
            "color": "#10b981",
            "shape": "square",
            "text": "E",
            "id": m_id
        })
        marker_details[m_id] = {
            "type": "Earnings",
            "date": c_date,
            "session": str(earn.get("time", "AMC")).upper(),
            "eps_actual": earn.get("eps_actual") if earn.get("eps_actual") is not None else (extra_info.get("eps_actual") if extra_info else None),
            "eps_estimate": earn.get("eps_estimate") if earn.get("eps_estimate") is not None else (extra_info.get("eps_estimate") if extra_info else None),
            "eps_surprise_pct": earn.get("eps_surprise_pct") if earn.get("eps_surprise_pct") is not None else (extra_info.get("eps_surprise_pct") if extra_info else None),
            "revenue_actual": earn.get("revenue_actual") if earn.get("revenue_actual") is not None else (extra_info.get("revenue_actual") if extra_info else None),
            "revenue_estimate": earn.get("revenue_estimate") if earn.get("revenue_estimate") is not None else (extra_info.get("revenue_estimate") if extra_info else None)
        }

    # Fallback to PostgreSQL reports if FMP has few/no earnings (e.g. non-AAPL tickers)
    db_reports_df = get_ticker_reports(ticker)
    if not db_reports_df.empty:
        for _, r in db_reports_df.iterrows():
            raw_date = str(r["date"])
            c_date = find_nearest_candle_date(raw_date)
            if not c_date and r.get("move_date"):
                c_date = find_nearest_candle_date(str(r["move_date"]))
            if not c_date or c_date in seen_earnings_dates:
                continue
            seen_earnings_dates.add(c_date)
            m_id = f"e_{c_date}"
            time_str = str(r.get("time", "")).lower()
            session_name = "AMC" if "16:" in time_str or "amc" in time_str or "after" in time_str else "BMO"
            quarter_name = f"{r.get('fiscal_period', '')} {r.get('fiscal_year', '')}".strip()
            extra_info = find_matching_earnings_extra(extra_map, raw_date)
            if not extra_info and r.get("move_date"):
                extra_info = find_matching_earnings_extra(extra_map, str(r["move_date"]))

            markers.append({
                "time": c_date,
                "position": "belowBar",
                "color": "#10b981",
                "shape": "square",
                "text": "E",
                "id": m_id
            })
            marker_details[m_id] = {
                "type": "Earnings",
                "date": c_date,
                "session": (extra_info.get("time") if extra_info else None) or session_name,
                "quarter": quarter_name,
                "eps_actual": extra_info.get("eps_actual") if extra_info else None,
                "eps_estimate": extra_info.get("eps_estimate") if extra_info else None,
                "eps_surprise_pct": extra_info.get("eps_surprise_pct") if extra_info else None,
                "revenue_actual": extra_info.get("revenue_actual") if extra_info else None,
                "revenue_estimate": extra_info.get("revenue_estimate") if extra_info else None
            }

    # B. Dividends Markers ('D')
    seen_dividend_dates = set()
    div_list = fmp_client.get_dividends_history(ticker, limit=30)
    if not div_list:
        div_list = polygon_client.fetch_dividends(ticker)
    for div in div_list:
        raw_date = div.get("ex_date") or div.get("date")
        if not raw_date:
            continue
        c_date = find_nearest_candle_date(raw_date)
        if not c_date or c_date in seen_dividend_dates:
            continue
        seen_dividend_dates.add(c_date)
        m_id = f"d_{c_date}"
        markers.append({
            "time": c_date,
            "position": "belowBar",
            "color": "#d97706",
            "shape": "circle",
            "text": "D",
            "id": m_id
        })
        marker_details[m_id] = {
            "type": "Dividend",
            "ex_date": c_date,
            "amount": div.get("amount") or div.get("dividend"),
            "pay_date": div.get("pay_date") or div.get("paymentDate"),
            "yield": div.get("yield"),
            "frequency": div.get("frequency", "Quarterly")
        }

    return {
        "ticker": ticker,
        "candles": candles,
        "markers": markers,
        "marker_details": marker_details
    }


@app.get("/api/ticker/{ticker}/chart/intraday")
async def api_intraday_chart(
    ticker: str,
    date: Optional[str] = None,
    multiplier: int = Query(5, ge=1, le=60),
    days_back: int = Query(30, ge=1, le=90)
):
    """
    Returns continuous multi-day 5m (or 1m) intraday candlestick series around the requested date,
    allowing full scrolling to yesterday, day before, today, and future days.
    """
    ticker = ticker.strip().upper()
    today_dt = datetime.date.today()
    if date:
        try:
            req_dt = datetime.date.fromisoformat(date)
        except Exception:
            req_dt = today_dt
    else:
        req_dt = today_dt

    # 30 days window around requested date so user can freely scroll backwards & forwards
    start_date = (req_dt - datetime.timedelta(days=days_back)).isoformat()
    end_date = min(today_dt, req_dt + datetime.timedelta(days=10)).isoformat()

    # 1. Try DB
    df = get_intraday_chart_data(ticker, start_date=start_date, end_date=end_date)
    candles = []

    if not df.empty:
        # Sort chronologically by date and time
        df = df.sort_values(["date", "time"])
        for _, r in df.iterrows():
            dt_str = f"{r['date']} {r['time']}"
            try:
                dt_obj = datetime.datetime.strptime(dt_str, "%Y-%m-%d %H:%M:%S")
                ts = int(dt_obj.replace(tzinfo=datetime.timezone.utc).timestamp())
            except Exception:
                ts = None
            candles.append({
                "time": dt_str,
                "timestamp": ts,
                "date": str(r["date"]),
                "time_str": r["time_str"],
                "open": float(r["open"]),
                "high": float(r["high"]),
                "low": float(r["low"]),
                "close": float(r["close"]),
                "volume": float(r["volume"]),
                "session_type": r.get("session_type", "regular")
            })
    else:
        # Fallback to Polygon
        poly_bars = polygon_client.fetch_intraday_bars(
            ticker,
            start_date=start_date,
            end_date=end_date,
            multiplier=multiplier,
            timespan="minute"
        )
        for b in poly_bars:
            candles.append({
                "time": b["datetime_et"],
                "timestamp": b.get("timestamp"),
                "date": b.get("date"),
                "time_str": b["time_str"],
                "open": b["open"],
                "high": b["high"],
                "low": b["low"],
                "close": b["close"],
                "volume": b["volume"],
                "session_type": b["session_type"]
            })

    return {
        "ticker": ticker,
        "date": req_dt.isoformat(),
        "start_date": start_date,
        "end_date": end_date,
        "multiplier": multiplier,
        "candles": candles,
        "count": len(candles)
    }


# ============================================================
# EARNINGS REACTION MATRIX (8 QUARTERS)
# ============================================================

@app.get("/api/ticker/{ticker}/earnings-moves")
async def api_earnings_moves(
    ticker: str,
    window: str = Query("04:00_09:30", pattern="^(04:00_07:00|04:00_09:00|04:00_09:30)$"),
    limit: int = Query(8, ge=1, le=16)
):
    """
    Returns the 8-quarter Earnings Reaction Matrix with 1-D postmarket %, 0-D premarket % (interval filtered), and intraday metrics.
    """
    ticker = ticker.strip().upper()
    matrix = earnings_reaction_client.get_earnings_reaction_matrix(
        ticker=ticker,
        limit=limit,
        premarket_window=window
    )
    return {
        "ticker": ticker,
        "premarket_window": window,
        "reports_count": len(matrix),
        "matrix": matrix
    }


# ============================================================
# EARNINGS REACTION MATRIX EXPORT & CORRELATIONS
# ============================================================

METRIC_LABELS = {
    "close_to_close": "Весь день (Close-to-Close %)",
    "premarket": "Премаркет %",
    "intraday": "Интрадей (Open → Close %)",
    "postmarket": "Постмаркет %"
}

@app.get("/api/ticker/{ticker}/earnings-moves/export")
async def export_earnings_moves(
    ticker: str,
    r_minus_1_metric: str = Query("close_to_close"),
    r_0_metric: str = Query("close_to_close"),
    r_plus_1_metric: str = Query("close_to_close"),
    r_plus_2_metric: str = Query("close_to_close")
):
    """
    Exports the 8-quarter Earnings Reaction Matrix to a styled Excel spreadsheet with selected metrics.
    """
    ticker = ticker.strip().upper()
    matrix = earnings_reaction_client.get_earnings_reaction_matrix(ticker=ticker, limit=8)
    if not matrix:
        raise HTTPException(status_code=400, detail="Нет данных по отчетам для экспорта.")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{ticker} Earnings Matrix"

    # Title row
    ws.merge_cells("A1:L1")
    title_cell = ws["A1"]
    title_cell.value = f"CHEPUH / FinWhale • Матрица движений на отчетности: {ticker} (Последние 8 кварталов)"
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="1E293B")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 28

    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    border_thin = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

    pos_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
    pos_font = Font(name="Calibri", size=10, color="1B5E20", bold=True)
    neg_fill = PatternFill(start_color="FFEBEE", end_color="FFEBEE", fill_type="solid")
    neg_font = Font(name="Calibri", size=10, color="B71C1C", bold=True)

    col_metrics = {
        "R-1": r_minus_1_metric,
        "R": r_0_metric,
        "R+1": r_plus_1_metric,
        "R+2": r_plus_2_metric
    }

    headers = [
        "Квартал",
        "Дата отчета",
        "Сессия",
        "EPS Факт ($)",
        "EPS Ожид. ($)",
        "Сюрприз EPS (%)",
        "Выручка Факт",
        "Выручка Ожид.",
        f"R-1 ({METRIC_LABELS.get(r_minus_1_metric, r_minus_1_metric)})",
        f"R ({METRIC_LABELS.get(r_0_metric, r_0_metric)})",
        f"R+1 ({METRIC_LABELS.get(r_plus_1_metric, r_plus_1_metric)})",
        f"R+2 ({METRIC_LABELS.get(r_plus_2_metric, r_plus_2_metric)})"
    ]

    ws.append(headers)
    ws.row_dimensions[2].height = 26

    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=2, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = cell_border

    # Day metric values collector for summary stats
    day_values = {"R-1": [], "R": [], "R+1": [], "R+2": []}

    for row_idx, r in enumerate(matrix, start=3):
        days_dict = r.get("days", {})
        row_cells = [
            r.get("quarter", "—"),
            r.get("date", "—"),
            r.get("session", "—"),
            r.get("eps_actual"),
            r.get("eps_estimate"),
            (f"{r.get('eps_surprise_pct'):+.2f}%" if r.get("eps_surprise_pct") is not None else "—"),
            r.get("revenue_actual"),
            r.get("revenue_estimate")
        ]

        for d_key in ["R-1", "R", "R+1", "R+2"]:
            m_name = col_metrics[d_key]
            d_info = days_dict.get(d_key, {})
            val = d_info.get(m_name)
            if val is not None and isinstance(val, (int, float)):
                day_values[d_key].append(float(val))
                row_cells.append(f"{val:+.2f}%")
            else:
                row_cells.append("—")

        ws.append(row_cells)
        ws.row_dimensions[row_idx].height = 20

        # Style data row
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.border = cell_border
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_idx in [9, 10, 11, 12]:
                val_str = str(cell.value or "")
                if val_str.startswith("+"):
                    cell.fill = pos_fill
                    cell.font = pos_font
                elif val_str.startswith("-"):
                    cell.fill = neg_fill
                    cell.font = neg_font

    # Summary Statistics rows (Mean, Median, Max, Min, Win Rate)
    def calc_stats(vals):
        if not vals:
            return "—", "—", "—", "—", "—"
        mean_v = np.mean(vals)
        med_v = np.median(vals)
        max_v = np.max(vals)
        min_v = np.min(vals)
        win_rate = (sum(1 for v in vals if v > 0) / len(vals)) * 100
        return f"{mean_v:+.2f}%", f"{med_v:+.2f}%", f"{max_v:+.2f}%", f"{min_v:+.2f}%", f"{win_rate:.1f}%"

    stat_rows = [
        ("Среднее (Mean)", 0),
        ("Медиана (Median)", 1),
        ("Максимум (Max)", 2),
        ("Минимум (Min)", 3),
        ("% Положительных (Win Rate)", 4)
    ]

    stat_start = len(matrix) + 4
    for s_idx, (stat_name, stat_pos) in enumerate(stat_rows, start=stat_start):
        s_row = [stat_name, "", "", "", "", "", "", ""]
        for d_key in ["R-1", "R", "R+1", "R+2"]:
            stats_tuple = calc_stats(day_values[d_key])
            s_row.append(stats_tuple[stat_pos])
        ws.append(s_row)
        ws.merge_cells(start_row=s_idx, start_column=1, end_row=s_idx, end_column=8)
        c_label = ws.cell(row=s_idx, column=1)
        c_label.font = Font(name="Calibri", size=9.5, bold=True, color="475569")
        c_label.alignment = Alignment(horizontal="right", vertical="center")
        c_label.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
        for c_i in range(1, 13):
            ws.cell(row=s_idx, column=c_i).border = cell_border
            if c_i >= 9:
                ws.cell(row=s_idx, column=c_i).alignment = Alignment(horizontal="center", vertical="center")
                ws.cell(row=s_idx, column=c_i).font = Font(name="Calibri", size=9.5, bold=True)

    # Column dimensions
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 13)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"{ticker}_Earnings_Moves_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


# ============================================================
# QUANT CORRELATIONS & BETA MATRIX
# ============================================================

@app.get("/api/ticker/{ticker}/correlations")
async def api_correlations(
    ticker: str,
    period: str = Query("1y", pattern="^(2w|1m|3m|6m|ytd|1y|earnings_8q)$"),
    peers: Optional[str] = None,
    benchmarks: Optional[str] = None
):
    """
    Calculates Pearson correlation and Beta coefficients against SPY, QQQ, SOXX, IGV, DRAM, and stock peers.
    Supports period="earnings_8q" for calculations restricted to earnings dates.
    """
    ticker = ticker.strip().upper()
    peer_list = [p.strip().upper() for p in peers.split(",") if p.strip()] if peers is not None else None
    bm_list = [b.strip().upper() for b in benchmarks.split(",") if b.strip()] if benchmarks is not None else None
    result = quant_client.calculate_ticker_correlations(
        ticker=ticker,
        period=period,
        peers=peer_list,
        benchmarks=bm_list
    )
    return result


# ============================================================
# DIVIDENDS HISTORY
# ============================================================

@app.get("/api/ticker/{ticker}/dividends")
async def api_dividends(ticker: str):
    ticker = ticker.strip().upper()
    db_divs = get_ticker_dividends_db(ticker)
    if not db_divs:
        # 1. Try FMP
        fmp_divs = fmp_client.get_dividends_history(ticker, limit=24)
        if not fmp_divs:
            # 2. Try Polygon fallback
            fmp_divs = polygon_client.fetch_dividends(ticker, limit=24)
        if fmp_divs:
            save_ticker_dividends_db(ticker, fmp_divs)
            db_divs = fmp_divs

    return {
        "ticker": ticker,
        "count": len(db_divs),
        "dividends": db_divs
    }


# ============================================================
# ON-DEMAND SYNC
# ============================================================

@app.post("/api/ticker/{ticker}/sync")
async def api_sync_ticker(ticker: str):
    ticker = ticker.strip().upper()
    # 1. Sync Profile & Peers
    prof = fmp_client.get_company_profile(ticker)
    peers = fmp_client.get_stock_peers(ticker)
    prof["peers"] = [p["symbol"] for p in peers if p.get("symbol")]
    save_company_profile_db(ticker, prof)

    # 2. Sync Dividends
    divs = fmp_client.get_dividends_history(ticker, limit=24)
    if divs:
        save_ticker_dividends_db(ticker, divs)

    # 3. Fetch latest Daily bars
    daily_bars = polygon_client.fetch_daily_bars(ticker, start_date="2023-01-01")

    return {
        "ticker": ticker,
        "synced": True,
        "daily_bars": len(daily_bars),
        "dividends": len(divs),
        "peers": prof["peers"]
    }


# ============================================================
# CORRELATION & BETA CALCULATOR ENDPOINTS
# ============================================================

def clean_value(value):
    if value is None:
        return None
    if isinstance(value, (np.integer, int)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        if np.isnan(value) or not np.isfinite(value):
            return None
        return float(value)
    return value


def clean_dict(data):
    if isinstance(data, dict):
        return {str(k): clean_dict(v) for k, v in data.items()}
    if isinstance(data, list):
        return [clean_dict(v) for v in data]
    return clean_value(data)


@app.post("/api/calculate")
async def api_calculate_single(request: Request):
    try:
        data = await request.json()
        if not data:
            raise HTTPException(status_code=400, detail="Пустой запрос.")

        stock = str(data.get("stock", "")).strip().upper()
        benchmark = str(data.get("benchmark", "")).strip().upper()
        benchmark_move = float(data.get("benchmark_move", 0.01))
        asset_move = data.get("asset_move")
        if asset_move is not None:
            asset_move = float(asset_move)

        if not stock or not benchmark:
            raise HTTPException(status_code=400, detail="Не указан тикер или benchmark.")
        if stock == benchmark:
            raise HTTPException(status_code=400, detail="Акции должны быть разными.")

        df = get_stock_data(stock, benchmark)
        if df is None or df.empty:
            raise HTTPException(status_code=404, detail="Данные не найдены.")

        results, _ = calculate_all_periods(
            df,
            stock,
            benchmark,
            benchmark_move=benchmark_move,
            asset_move=asset_move
        )

        return {
            "stock": stock,
            "benchmark": benchmark,
            "benchmark_move": benchmark_move,
            "asset_move": asset_move,
            "results": clean_dict(results),
            "chart": []
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/calculate-groups")
@app.post("/api/calculate_groups")
async def api_calculate_groups(request: Request):
    try:
        data = await request.json()
        if not data:
            raise HTTPException(status_code=400, detail="Пустой запрос.")

        groups = data.get("groups")
        benchmark_raw = data.get("benchmarks") or data.get("benchmark") or data.get("reference")
        period = data.get("period") or "1y"
        benchmark_move = float(data.get("benchmark_move", 0.01))
        benchmark_moves = data.get("benchmark_moves", {})
        asset_moves = data.get("asset_moves", {})

        if not groups or not isinstance(groups, dict):
            raise HTTPException(status_code=400, detail="Добавьте хотя бы одну группу с тикерами.")
        if not benchmark_raw:
            raise HTTPException(status_code=400, detail="Не указан Market Asset (Benchmark).")

        benchmarks_list = []
        if isinstance(benchmark_raw, (list, tuple, set)):
            benchmarks_list = [str(b).strip().upper() for b in benchmark_raw if str(b).strip()]
        elif isinstance(benchmark_raw, str):
            benchmarks_list = [b.strip().upper() for b in re.split(r"[\s,;]+", benchmark_raw) if b.strip()]

        clean_benchmarks = []
        for bm in benchmarks_list:
            if bm and bm not in clean_benchmarks:
                clean_benchmarks.append(bm)

        if not clean_benchmarks:
            raise HTTPException(status_code=400, detail="Не указан валидный Market Asset (Benchmark).")

        cleaned_groups = {}
        all_requested_tickers = []
        for group_name, tickers in groups.items():
            if not isinstance(tickers, list):
                continue
            clean_tickers = []
            for ticker in tickers:
                if not isinstance(ticker, str):
                    continue
                ticker = ticker.strip().upper()
                if ticker and ticker not in clean_benchmarks and ticker not in clean_tickers:
                    clean_tickers.append(ticker)
                    if ticker not in all_requested_tickers:
                        all_requested_tickers.append(ticker)
            if clean_tickers:
                cleaned_groups[str(group_name).strip() or "Group"] = clean_tickers

        if not cleaned_groups or not all_requested_tickers:
            raise HTTPException(status_code=400, detail="Нет валидных тикеров для анализа.")

        all_symbols = list(dict.fromkeys(all_requested_tickers + clean_benchmarks))
        df = get_stock_data(*all_symbols)
        if df is None or df.empty:
            raise HTTPException(status_code=404, detail="Данные для указанных тикеров не найдены в базе данных.")

        available_tickers = set(df["ticker"].astype(str).str.strip().str.upper().unique())
        missing_benchmarks = [bm for bm in clean_benchmarks if bm not in available_tickers]
        if missing_benchmarks:
            raise HTTPException(status_code=404, detail=f"Market Asset '{', '.join(missing_benchmarks)}' не найден в базе данных.")

        missing_tickers = [t for t in all_requested_tickers if t not in available_tickers]
        active_groups = {}
        for g_name, t_list in cleaned_groups.items():
            valid_t = [t for t in t_list if t in available_tickers]
            if valid_t:
                active_groups[g_name] = valid_t

        if not active_groups:
            raise HTTPException(status_code=404, detail="Ни один из указанных тикеров не найден в базе данных.")

        cleaned_asset_moves = {}
        if isinstance(asset_moves, dict):
            for t, m in asset_moves.items():
                try:
                    m = float(m)
                    if np.isfinite(m):
                        cleaned_asset_moves[str(t).strip().upper()] = m
                except (TypeError, ValueError):
                    continue

        analysis = calculate_groups_analysis(
            df=df,
            groups=active_groups,
            benchmarks=clean_benchmarks,
            period=period,
            benchmark_move=benchmark_move,
            benchmark_moves=benchmark_moves if isinstance(benchmark_moves, dict) else None,
            asset_moves=cleaned_asset_moves
        )

        ignored_tickers = analysis.get("ignored_tickers", {})
        partial_tickers = analysis.get("partial_tickers", {})
        warning_parts = []
        if missing_tickers:
            warning_parts.append(
                f"Пропущено {len(missing_tickers)} тикеров (нет в БД): {', '.join(missing_tickers[:8])}"
                + (f" и ещё {len(missing_tickers) - 8}..." if len(missing_tickers) > 8 else "")
            )
        if ignored_tickers:
            ignored_list = list(ignored_tickers.keys()) if isinstance(ignored_tickers, dict) else list(ignored_tickers)
            warning_parts.append(
                f"Не удалось рассчитать {len(ignored_list)} тикеров (< 2 наблюдений): {', '.join(ignored_list[:8])}"
                + (f" и ещё {len(ignored_list) - 8}..." if len(ignored_list) > 8 else "")
            )
        if partial_tickers:
            partial_list = list(partial_tickers.keys())
            warning_parts.append(
                f"{len(partial_list)} тикеров имеют историю меньше периода {period}: {', '.join(partial_list[:6])}"
                + (f" и ещё {len(partial_list) - 6}..." if len(partial_list) > 6 else "")
            )

        warning_message = " | ".join(warning_parts) if warning_parts else None

        return {
            "groups": active_groups,
            "benchmarks": clean_benchmarks,
            "primary_benchmark": clean_benchmarks[0],
            "benchmark": clean_benchmarks[0],
            "reference": clean_benchmarks[0],
            "benchmark_moves": clean_dict(analysis.get("benchmark_moves", {})),
            "benchmark_move": benchmark_move,
            "asset_moves": cleaned_asset_moves,
            "tickers": analysis.get("tickers", []),
            "all_symbols": analysis.get("all_symbols", []),
            "period": analysis.get("period", period),
            "period_start": analysis.get("period_start"),
            "period_end": analysis.get("period_end"),
            "correlation": clean_dict(analysis.get("correlation", {})),
            "beta": clean_dict(analysis.get("beta", {})),
            "benchmark_beta": clean_dict(analysis.get("benchmark_beta", {})),
            "expected_range": clean_dict(analysis.get("expected_range", {})),
            "benchmarks_data": clean_dict(analysis.get("benchmarks_data", {})),
            "group_summary": clean_dict(analysis.get("group_summary", {})),
            "missing_tickers": missing_tickers,
            "ignored_tickers": clean_dict(ignored_tickers),
            "partial_tickers": clean_dict(partial_tickers),
            "warning": warning_message,
            "total_requested": len(all_requested_tickers),
            "total_analyzed": len(analysis.get("tickers", []))
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))


@app.post("/api/export-excel")
async def api_export_excel(request: Request):
    try:
        from fastapi.responses import Response
        data = await request.json() or {}

        rows = data.get("rows", [])
        benchmarks = data.get("benchmarks", [])
        benchmark = str(data.get("benchmark", "")).strip().upper() or "BENCHMARK"
        period = str(data.get("period", "1y"))
        is_multi = bool(data.get("is_multi_benchmark", False) or (isinstance(benchmarks, list) and len(benchmarks) > 1))

        if not rows:
            raise HTTPException(status_code=400, detail="Нет данных для экспорта.")

        if not benchmarks:
            benchmarks = [benchmark]

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Analysis Results"

        header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        border_thin = Side(border_style="thin", color="E2E8F0")
        cell_border = Border(left=border_thin, right=border_thin, top=border_thin, bottom=border_thin)

        if is_multi and len(benchmarks) > 1:
            headers = ["Ticker"]
            for bm in benchmarks:
                headers.extend([f"Corr ({bm})", f"Beta ({bm})", f"Exp Move ({bm})"])
            headers.extend(["Actual Move", "Daily Vol", "Observations", "Period"])
            ws.append(headers)

            for r in rows:
                ticker = str(r.get("ticker", "")).strip().upper()
                row_data = [ticker]
                for bm in benchmarks:
                    row_data.append(clean_value(r.get(f"correlation_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("correlation")))
                    row_data.append(clean_value(r.get(f"beta_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("beta")))
                    row_data.append(clean_value(r.get(f"expected_move_{bm}") or r.get("benchmarks_metrics", {}).get(bm, {}).get("expected_move")))
                row_data.extend([clean_value(r.get("actual_move")), clean_value(r.get("volatility")), clean_value(r.get("observations")), period])
                ws.append(row_data)
        else:
            headers = ["Ticker", f"Corr ({benchmark})", f"Beta ({benchmark})", f"Exp Move ({benchmark})", "Actual Move", "Difference", "Normalized", "Daily Vol", "Residual Vol", "Observations"]
            ws.append(headers)
            for r in rows:
                ws.append([
                    str(r.get("ticker", "")).strip().upper(),
                    clean_value(r.get("correlation")),
                    clean_value(r.get("beta")),
                    clean_value(r.get("expected_move")),
                    clean_value(r.get("actual_move")),
                    clean_value(r.get("difference")),
                    clean_value(r.get("normalized")),
                    clean_value(r.get("volatility")),
                    clean_value(r.get("residual_volatility")),
                    clean_value(r.get("observations"))
                ])

        ws.row_dimensions[1].height = 24
        for col in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment
            cell.border = cell_border

        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)

        filename = f"Quant_Analysis_{period}_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        return Response(
            content=output.getvalue(),
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f'attachment; filename="{filename}"'}
        )
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))



if __name__ == "__main__":
    import uvicorn
    import socket

    def find_free_port(preferred_port=8000):
        for port in [preferred_port, 8080, 8001, 8002, 5000, 5050]:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                try:
                    s.bind(("127.0.0.1", port))
                    return port
                except OSError:
                    continue
        return preferred_port

    target_port = int(os.getenv("PORT", find_free_port(8000)))
    print("\n" + "=" * 55)
    print(f">> Helion Quant Platform: http://127.0.0.1:{target_port}")
    print(f">> Swagger API Docs:      http://127.0.0.1:{target_port}/docs")
    print("=" * 55 + "\n")
    uvicorn.run("app_fastapi:app", host="127.0.0.1", port=target_port, reload=True)
